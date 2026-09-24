"""
Массовая привязка дублей артикулов (alias_of).
"""
from django.core.management.base import BaseCommand
from django.db import transaction

from inventory.models import Product


def _find_ultimate_main(product):
    seen = set()
    current = product
    while current.alias_of_id:
        if current.id in seen:
            break
        seen.add(current.id)
        current = current.alias_of
    return current


class Command(BaseCommand):
    help = "Массовая привязка дублей артикулов"

    def add_arguments(self, parser):
        parser.add_argument("--auto", type=str, help="Путь к xlsx")
        parser.add_argument("--csv", type=str, help="Путь к CSV")
        parser.add_argument("--sheet", type=str, default=None)
        parser.add_argument("--dry-run", action="store_true")
        parser.add_argument("--create-missing", action="store_true")

    def handle(self, *args, **options):
        if options["auto"]:
            pairs = self._read_xlsx(options["auto"], options["sheet"])
        elif options["csv"]:
            pairs = self._read_csv(options["csv"])
        else:
            self.stderr.write("Укажи --auto <xlsx> или --csv <file>")
            return

        # TWO-PHASE: разрешаем конфликты и циклы ДО применения
        self.stdout.write(f"Найдено пар (до нормализации): {len(pairs)}")

        # CYCLE-V3: граф-based нормализация
        # 1. Собираем все пары
        # 2. Убираем циклы (прямые и транзитивные) через DFS
        # 3. Убираем дубликаты "alias имеет два main"
        # 4. Убираем "main который уже сам alias"

        from collections import defaultdict

        # candidates: list of (main, alias) в порядке появления
        candidates = []
        for main_art, alias_art in pairs:
            if not main_art or not alias_art or main_art == alias_art:
                continue
            candidates.append((main_art, alias_art))

        # Граф: alias_art → main_art (то, что уже приняли)
        chosen_map = {}          # alias → main
        conflicts = []           # (alias, why, main1, main2)

        def _path_exists(start, target, mapping):
            """Есть ли путь от start до target по mapping."""
            if start == target:
                return True
            seen = set()
            stack = [start]
            while stack:
                cur = stack.pop()
                if cur in seen:
                    continue
                seen.add(cur)
                nxt = mapping.get(cur)
                if nxt is None:
                    continue
                if nxt == target:
                    return True
                stack.append(nxt)
            return False

        for main_art, alias_art in candidates:
            # Проверка 1: main_art уже является alias (значит он не может быть главным)
            if main_art in chosen_map:
                conflicts.append((alias_art, "main уже alias",
                                  main_art, chosen_map[main_art]))
                continue

            # Проверка 2: alias_art уже привязан к другому main
            existing = chosen_map.get(alias_art)
            if existing is not None and existing != main_art:
                conflicts.append((alias_art, "alias уже привязан",
                                  existing, main_art))
                continue

            # Проверка 3: цикл. Если от main_art есть путь до alias_art —
            # значит добавить alias_art → main_art нельзя (создаст цикл).
            # Временная карта: chosen_map + новое ребро
            test_map = dict(chosen_map)
            test_map[alias_art] = main_art
            # Ищем путь от main_art до alias_art в test_map.
            # Если нашли — цикл.
            if _path_exists(main_art, alias_art, test_map):
                conflicts.append((alias_art, "цикл",
                                  main_art, alias_art))
                continue

            # Принимаем
            chosen_map[alias_art] = main_art

        # Пересобираем нормализованные пары
        normalized_pairs = [(m, a) for a, m in chosen_map.items()]

        pairs = normalized_pairs
        self.stdout.write(f"Найдено пар (после нормализации): {len(pairs)}")
        if conflicts:
            self.stdout.write(f"⚠️  Разрешено конфликтов: {len(conflicts)}")
            for c in conflicts[:15]:
                a, why, m1, m2 = c
                self.stdout.write(f"     {a}: {why} ({m1} vs {m2})")
        self.stdout.write("")

        dry = options["dry_run"]
        create_missing = options["create_missing"]

        created_products = []
        linked = 0
        skipped_missing = 0
        skipped_same = 0
        skipped_exists = 0
        errors = 0
        processed = set()

        for main_art, alias_art in pairs:
            if not main_art or not alias_art:
                skipped_missing += 1
                continue
            if main_art == alias_art:
                skipped_same += 1
                continue

            main_p = Product.objects.filter(article__iexact=main_art).first()
            alias_p = Product.objects.filter(article__iexact=alias_art).first()

            # Умное создание: копируем поля от существующего парного
            will_create_main = False
            will_create_alias = False

            if main_p is None and create_missing:
                if dry:
                    will_create_main = True
                    self.stdout.write(f"  🆕 будет создан main={main_art} (копия от {alias_art})")
                else:
                    main_p = Product.objects.create(
                        article=main_art,
                        name=(alias_p.name if alias_p else f"Артикул {main_art}"),
                        product_type=(alias_p.product_type if alias_p else "part"),
                        uom=(alias_p.uom if alias_p else "pcs"),
                        weight_g=(alias_p.weight_g if alias_p else 0),
                        is_active=True,
                    )
                    created_products.append(main_art)
            if alias_p is None and create_missing:
                if dry:
                    will_create_alias = True
                    self.stdout.write(f"  🆕 будет создан alias={alias_art} (копия от {main_art})")
                else:
                    alias_p = Product.objects.create(
                        article=alias_art,
                        name=(main_p.name if main_p else f"Артикул {alias_art}"),
                        product_type=(main_p.product_type if main_p else "part"),
                        uom=(main_p.uom if main_p else "pcs"),
                        weight_g=(main_p.weight_g if main_p else 0),
                        is_active=False,  # старый скрываем
                    )
                    created_products.append(alias_art)

            # В dry-run не можем проверить связи для несуществующих объектов
            if dry and (will_create_main or will_create_alias):
                linked += 1
                continue

            # Если после попыток создания кто-то всё ещё None (без --create-missing)
            if not create_missing and (main_p is None or alias_p is None):
                missing = []
                if main_p is None:
                    missing.append(f"main={main_art}")
                if alias_p is None:
                    missing.append(f"alias={alias_art}")
                if dry and skipped_missing < 15:
                    self.stdout.write(f"  ⏭  нет в БД: {', '.join(missing)}")
                skipped_missing += 1
                continue

            if alias_p.alias_of_id == main_p.id:
                skipped_exists += 1
                continue

            ultimate_main = _find_ultimate_main(main_p)
            if ultimate_main.id == alias_p.id:
                self.stdout.write(f"  ❌ Цикл: {alias_art} → {main_art}")
                errors += 1
                continue

            key = (ultimate_main.id, alias_p.id)
            if key in processed:
                continue
            processed.add(key)

            if dry:
                if linked < 25:
                    self.stdout.write(f"  📝 {alias_art} → {ultimate_main.article}")
            else:
                with transaction.atomic():
                    for child in alias_p.aliases.all():
                        child.alias_of = ultimate_main
                        child.save(update_fields=["alias_of", "updated_at"])
                    alias_p.alias_of = ultimate_main
                    alias_p.save(update_fields=["alias_of", "updated_at"])
                self.stdout.write(f"  ✅ {alias_art} → {ultimate_main.article}")
            linked += 1

        self.stdout.write("")
        self.stdout.write("═" * 60)
        self.stdout.write(f"  Связано:          {linked}")
        self.stdout.write(f"  Уже связано:      {skipped_exists}")
        self.stdout.write(f"  Одинаковые:       {skipped_same}")
        self.stdout.write(f"  Нет в БД:         {skipped_missing}")
        self.stdout.write(f"  Ошибок:           {errors}")
        if created_products:
            self.stdout.write(f"  Создано (пустых): {len(created_products)}")
        if dry:
            self.stdout.write("  ⚠️  DRY RUN — ничего не сохранено")
        self.stdout.write("═" * 60)

    def _read_xlsx(self, path, sheet_name):
        try:
            from openpyxl import load_workbook
        except ImportError:
            self.stderr.write("openpyxl не установлен")
            return []
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[sheet_name] if sheet_name else wb.active
        self.stdout.write(f"Лист: {ws.title}")
        # Колонки xlsx:
        #   A (0) = № п/п          — пропускаем
        #   B (1) = старый литьё
        #   C (2) = новый литьё
        #   D (3) = старый мех
        #   E (4) = новый мех
        pairs = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or len(row) < 5:
                continue
            old_cast = self._clean(row[1])   # старый литьё
            new_cast = self._clean(row[2])   # новый литьё
            old_mech = self._clean(row[3])   # старый мех
            new_mech = self._clean(row[4])   # новый мех

            if old_cast and new_cast and old_cast != new_cast:
                pairs.append((new_cast, old_cast))  # main=новый, alias=старый
            if old_mech and new_mech and old_mech != new_mech:
                pairs.append((new_mech, old_mech))
        return pairs

    def _read_csv(self, path):
        import csv
        pairs = []
        with open(path, "r", encoding="utf-8-sig") as f:
            reader = csv.reader(f)
            for i, row in enumerate(reader):
                if len(row) < 2:
                    continue
                if i == 0 and not any(ch.isdigit() for ch in row[0]):
                    continue
                main = self._clean(row[0])
                alias = self._clean(row[1])
                if main and alias and main != alias:
                    pairs.append((main, alias))
        return pairs

    @staticmethod
    def _clean(v):
        if v is None:
            return ""
        s = str(v).strip()
        if s.lower() in ("none", "null", "nan", "—", "-", "нет"):
            return ""
        return s
