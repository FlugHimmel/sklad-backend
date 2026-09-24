"""
Генерация PDF из docx-шаблона через LibreOffice headless.

Использование:
    from warehouse.pdf_from_template import generate_from_docx_template
    pdf_bytes = generate_from_docx_template(
        template_path="/opt/sklad/templates/packing_vilo.docx",
        replacements={"NAME": "Колесо", "QTY": "50", ...},
    )

Требования:
    * LibreOffice установлен (soffice / libreoffice в PATH)
    * python-docx, pypdf в venv
"""
import os
import shutil
import subprocess
import tempfile
import uuid
from pathlib import Path

from docx import Document

# Импорт resolve из основного pdf-модуля (BOM + Mo-слоты)
from .pdf import _resolve_art_before  # noqa: E402

# Путь к soffice — жёстко прописанные варианты + поиск в PATH
_SOFFICE_CANDIDATES = [
    "/usr/bin/soffice",
    "/usr/bin/libreoffice",
    "/usr/local/bin/soffice",
    "/usr/local/bin/libreoffice",
    "/opt/libreoffice/program/soffice",
]

def _find_soffice():
    # Сначала — абсолютные пути
    for path in _SOFFICE_CANDIDATES:
        if os.path.exists(path) and os.access(path, os.X_OK):
            return path
    # Потом — PATH
    found = shutil.which("soffice") or shutil.which("libreoffice")
    if found:
        return found
    return None

_SOFFICE = _find_soffice()


MONTHS_RU = {
    1: "января", 2: "февраля", 3: "марта", 4: "апреля",
    5: "мая", 6: "июня", 7: "июля", 8: "августа",
    9: "сентября", 10: "октября", 11: "ноября", 12: "декабря",
}

MONTHS_RU_CAP = {
    1: "Января", 2: "Февраля", 3: "Марта", 4: "Апреля",
    5: "Мая", 6: "Июня", 7: "Июля", 8: "Августа",
    9: "Сентября", 10: "Октября", 11: "Ноября", 12: "Декабря",
}


def _replace_in_paragraph(paragraph, replacements: dict):
    """Заменяет плейсхолдеры в параграфе, склеивая runs.
    Placeholders должны быть без форматирования внутри (иначе нельзя гарантировать)."""
    full_text = "".join(r.text for r in paragraph.runs)
    if not full_text:
        return
    new_text = full_text
    for key, val in replacements.items():
        token = "{{" + key + "}}"
        if token in new_text:
            new_text = new_text.replace(token, str(val))
    if new_text != full_text and paragraph.runs:
        paragraph.runs[0].text = new_text
        for r in paragraph.runs[1:]:
            r.text = ""


def _replace_in_cell(cell, replacements: dict):
    for p in cell.paragraphs:
        _replace_in_paragraph(p, replacements)


def fill_docx_template(template_path: str, out_path: str, replacements: dict):
    """Открывает docx, меняет плейсхолдеры, сохраняет в out_path."""
    doc = Document(template_path)

    for p in doc.paragraphs:
        _replace_in_paragraph(p, replacements)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                _replace_in_cell(cell, replacements)

    # Также — в headers/footers если есть
    for section in doc.sections:
        for hf in (section.header, section.footer):
            if hf is None:
                continue
            for p in hf.paragraphs:
                _replace_in_paragraph(p, replacements)
            for table in hf.tables:
                for row in table.rows:
                    for cell in row.cells:
                        _replace_in_cell(cell, replacements)

    doc.save(out_path)


def _convert_docx_to_pdf(docx_path: str, out_dir: str) -> str:
    """Конвертирует docx в pdf через LibreOffice headless.
    Возвращает путь к получившемуся pdf.
    """
    if not _SOFFICE:
        raise RuntimeError(
            "LibreOffice не найден. Установи: apt install libreoffice-writer"
        )

    # Уникальный профиль пользователя — иначе soffice может зависнуть,
    # если другой экземпляр уже запущен.
    profile_dir = tempfile.mkdtemp(prefix="lo_profile_")
    try:
        cmd = [
            _SOFFICE,
            f"-env:UserInstallation=file://{profile_dir}",
            "--headless",
            "--norestore",
            "--convert-to", "pdf:writer_pdf_Export",
            "--outdir", out_dir,
            docx_path,
        ]

        # soffice — shell-скрипт, которому нужны dirname/basename/sed/uname.
        # У gunicorn PATH пустой — прописываем полный.
        env = os.environ.copy()
        env["PATH"] = "/usr/local/sbin:/usr/local/bin:/usr/sbin:/usr/bin:/sbin:/bin"
        env["HOME"] = profile_dir
        env["LANG"] = "ru_RU.UTF-8"
        env["LC_ALL"] = "ru_RU.UTF-8"

        proc = subprocess.run(
            cmd,
            env=env,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            timeout=60,
        )
        if proc.returncode != 0:
            raise RuntimeError(
                f"LibreOffice вернул код {proc.returncode}: "
                f"{proc.stderr.decode('utf-8', errors='ignore')[:500]}"
            )

        base = Path(docx_path).stem
        pdf_path = Path(out_dir) / f"{base}.pdf"
        if not pdf_path.exists():
            # Иногда имя с расширением дублируется
            candidates = list(Path(out_dir).glob(f"{base}*.pdf"))
            if candidates:
                pdf_path = candidates[0]
            else:
                raise RuntimeError(
                    f"LibreOffice не создал PDF в {out_dir}. "
                    f"stdout: {proc.stdout.decode('utf-8', errors='ignore')[:300]}"
                )
        return str(pdf_path)
    finally:
        shutil.rmtree(profile_dir, ignore_errors=True)


def generate_from_docx_template(
    template_path: str,
    replacements: dict,
) -> bytes:
    """Полный цикл: docx → заполненный docx → PDF. Возвращает bytes."""
    work = tempfile.mkdtemp(prefix="packing_")
    try:
        docx_filled = os.path.join(work, "filled.docx")
        fill_docx_template(template_path, docx_filled, replacements)
        pdf_path = _convert_docx_to_pdf(docx_filled, work)
        with open(pdf_path, "rb") as f:
            return f.read()
    finally:
        shutil.rmtree(work, ignore_errors=True)


def generate_bulk_from_docx_template(
    template_path: str,
    containers_data: list,
) -> bytes:
    """Несколько тар → один PDF (склеиваем страницы).
    containers_data = [{"replacements": {...}}, ...]
    """
    from pypdf import PdfWriter, PdfReader
    import io as _io

    writer = PdfWriter()
    for item in containers_data:
        pdf_bytes = generate_from_docx_template(
            template_path, item["replacements"],
        )
        reader = PdfReader(_io.BytesIO(pdf_bytes))
        for page in reader.pages:
            writer.add_page(page)

    out = _io.BytesIO()
    writer.write(out)
    return out.getvalue()


def build_replacements_for_container(container, product=None) -> dict:
    """Собирает dict replacements для одной тары.
    container — объект Container
    """
    from datetime import date
    from decimal import Decimal
    from warehouse.pdf import _resolve_art_before  # переиспользуем

    if product is None:
        product = container.product

    art_after = product.article if product else "—"
    name = product.name if product else "—"
    art_before = _resolve_art_before(product) or "—"

    qty = container.quantity or Decimal("0")
    wg = Decimal(str(product.weight_g)) if product and product.weight_g else Decimal("0")
    total_kg = (wg * Decimal(str(qty))) / Decimal("1000")

    today = date.today()

    return {
        "NAME": name,
        "ART_BEFORE": art_before,
        "ART_AFTER": art_after,
        "QTY": _fmt_qty(qty),
        "WEIGHT": str(total_kg.quantize(Decimal("0.001"))),
        "DAY": today.day,
        "MONTH": MONTHS_RU_CAP.get(today.month, ""),
        "YEAR": today.year,
    }


def _fmt_qty(q) -> str:
    from decimal import Decimal
    q = Decimal(str(q))
    if q == q.to_integral_value():
        return str(int(q))
    return f"{q:.3f}".rstrip("0").rstrip(".")


# ═════════════════════════════════════════════════════════════════════
# Многострочный режим: клонирование строк таблицы
# ═════════════════════════════════════════════════════════════════════
from docx.oxml.ns import qn as _qn


def _replace_in_tr(tr, replacements: dict):
    """Заменяет плейсхолдеры в XML-элементе <w:tr> (строка таблицы).
    Собирает текст из всех <w:r>/<w:t>, заменяет, пишет в первый <w:t>.
    """
    for p in tr.iter(_qn("w:p")):
        runs = list(p.findall(_qn("w:r")))
        if not runs:
            continue
        parts = []
        for r in runs:
            for t in r.findall(_qn("w:t")):
                parts.append(t.text or "")
        full = "".join(parts)
        new = full
        for k, v in replacements.items():
            new = new.replace("{{" + k + "}}", str(v))
        if new == full:
            continue
        written = False
        for r in runs:
            for t in r.findall(_qn("w:t")):
                if not written:
                    t.text = new
                    t.set(_qn("xml:space"), "preserve")
                    written = True
                else:
                    t.text = ""


def fill_docx_template_multi(template_path: str, out_path: str,
                              main_repl: dict, rows_data: list):
    """Открывает docx, меняет основные плейсхолдеры, размножает строку
    таблицы по rows_data (одна строка — один артикул).
    """
    from copy import deepcopy
    from docx.table import Table

    doc = Document(template_path)

    # Основные параграфы
    for p in doc.paragraphs:
        _replace_in_paragraph(p, main_repl)

    # Headers/footers
    for section in doc.sections:
        for hf in (section.header, section.footer):
            if hf is None:
                continue
            for p in hf.paragraphs:
                _replace_in_paragraph(p, main_repl)
            for table in hf.tables:
                for row in table.rows:
                    for cell in row.cells:
                        _replace_in_cell(cell, main_repl)

    # Таблицы: ищем строку-образец (обычно 2-я, после шапки)
    for table in doc.tables:
        if len(table.rows) < 2:
            for row in table.rows:
                for cell in row.cells:
                    _replace_in_cell(cell, main_repl)
            continue

        template_tr = table.rows[1]._tr
        parent = template_tr.getparent()

        # Чистый образец ДО заполнения (для клонирования)
        clean_tr = deepcopy(template_tr)

        # Заполняем первую строку данными #0
        _replace_in_tr(template_tr, {**main_repl, **rows_data[0]})

        # Клонируем остальные строки
        prev = template_tr
        for row_data in rows_data[1:]:
            new_tr = deepcopy(clean_tr)
            _replace_in_tr(new_tr, {**main_repl, **row_data})
            prev.addnext(new_tr)
            prev = new_tr

    doc.save(out_path)


def build_data_for_container(container):
    """Возвращает {"main": {...}, "rows": [{...}, ...]}.
    main — общие плейсхолдеры (дата, общий вес).
    rows — по одной записи на каждый артикул в таре.
    """
    from datetime import date
    from decimal import Decimal

    lines = list(container.lines.select_related("product").all())

    # Если в таре есть литьё — его артикул и есть «артикул до МО»
    # для всех деталей, даже если BOM указывает другое.
    casting_article_in_container = None
    for line in lines:
        prod = line.product
        if prod and prod.product_type == "casting":
            casting_article_in_container = prod.article
            break

    total_kg = Decimal("0")
    rows = []
    for idx, line in enumerate(lines, 1):
        p = line.product
        art_after = p.article if p else "—"
        name = p.name if p else "—"

        # 1. Литьё в таре → берём его (для деталей).
        # 2. Сама строка — литьё → «—».
        # 3. Иначе — BOM / Mo.
        if casting_article_in_container and p and p.product_type != "casting":
            art_before = casting_article_in_container
        elif p and p.product_type == "casting":
            art_before = "—"
        else:
            try:
                art_before = _resolve_art_before(p) if p else None
            except Exception:
                art_before = None
            art_before = art_before or "—"
        qty = line.quantity or Decimal("0")

        wg = Decimal(str(p.weight_g)) if p and p.weight_g else Decimal("0")
        line_kg = (wg * Decimal(str(qty))) / Decimal("1000")
        total_kg += line_kg

        rows.append({
            "SEQ": str(idx),
            "NAME": name,
            "ART_BEFORE": art_before,
            "ART_AFTER": art_after,
            "QTY": _fmt_qty(qty),
        })

    if not rows:
        rows.append({
            "SEQ": "1", "NAME": "—",
            "ART_BEFORE": "—", "ART_AFTER": "—", "QTY": "—",
        })

    today = date.today()
    return {
        "main": {
            "WEIGHT": str(total_kg.quantize(Decimal("0.001"))),
            "DAY": today.day,
            "MONTH": MONTHS_RU_CAP.get(today.month, ""),
            "YEAR": today.year,
        },
        "rows": rows,
    }


# ═════════════════════════════════════════════════════════════════════
# Обновлённые обёртки (принимают rows_data)
# ═════════════════════════════════════════════════════════════════════
_original_generate = generate_from_docx_template


def generate_from_docx_template_v2(template_path, replacements, rows_data=None):
    """Если rows_data передан — таблица размножается по строкам."""
    if rows_data is None:
        return _original_generate(template_path, replacements)
    work = tempfile.mkdtemp(prefix="packing_")
    try:
        docx_filled = os.path.join(work, "filled.docx")
        fill_docx_template_multi(
            template_path, docx_filled, replacements, rows_data,
        )
        pdf_path = _convert_docx_to_pdf(docx_filled, work)
        with open(pdf_path, "rb") as f:
            return f.read()
    finally:
        shutil.rmtree(work, ignore_errors=True)


# Переопределяем имя
generate_from_docx_template = generate_from_docx_template_v2


def generate_bulk_from_docx_template_v2(template_path, containers_data):
    """containers_data = [{"main": {...}, "rows": [...]}, ...]."""
    from pypdf import PdfWriter, PdfReader
    import io as _io

    writer = PdfWriter()
    for item in containers_data:
        pdf_bytes = generate_from_docx_template_v2(
            template_path, item["main"], rows_data=item["rows"],
        )
        reader = PdfReader(_io.BytesIO(pdf_bytes))
        for page in reader.pages:
            writer.add_page(page)

    out = _io.BytesIO()
    writer.write(out)
    return out.getvalue()


# Переопределяем bulk
generate_bulk_from_docx_template = generate_bulk_from_docx_template_v2

